"""Compara o FOR EMC38 do VIGIA-7S (registro.exportar_excel, openpyxl) com o do porte (registro_excel_ts.mjs).

    python registro_excel_cmp.py "<pasta do VIGIA>" reg.json saida_ts.xlsx [limite]

Gera o .xlsx de referência com o código original (só importa, não altera nada) e confere, aba por aba e célula
por célula: valor/fórmula, formato numérico, fonte, preenchimento, alinhamento, bordas, mesclagens, larguras de
coluna, formatação condicional e os gráficos (XML).
"""
import json
import os
import sys
import tempfile
import zipfile

vigia, arq_reg, arq_ts = sys.argv[1], sys.argv[2], sys.argv[3]
limite = float(sys.argv[4]) if len(sys.argv) > 4 else 1.0
sys.path.insert(0, os.path.join(vigia, "vigia7s"))
import registro  # noqa: E402
import openpyxl  # noqa: E402

reg = registro.completar(json.load(open(arq_reg, encoding="utf-8")))
arq_py = os.path.join(tempfile.mkdtemp(), "ref.xlsx")
registro.exportar_excel(reg, arq_py, limite)

wa, wb = openpyxl.load_workbook(arq_py), openpyxl.load_workbook(arq_ts)
dif = []


def d(onde, a, b):
    if a != b:
        dif.append(f"{onde}\n   py: {a!r}\n   ts: {b!r}")


def norm(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


d("abas", wa.sheetnames, wb.sheetnames)
n_cel = 0
for nome in wa.sheetnames:
    if nome not in wb.sheetnames:
        continue
    a, b = wa[nome], wb[nome]
    d(f"[{nome}] dimensão", a.max_row, b.max_row)
    d(f"[{nome}] mesclagens", sorted(map(str, a.merged_cells.ranges)), sorted(map(str, b.merged_cells.ranges)))
    def largura(ws, col):  # colunas agrupadas (min..max, como o exceljs grava) contam para cada coluna do grupo
        i = openpyxl.utils.column_index_from_string(col)
        for dim in ws.column_dimensions.values():
            if dim.min and dim.max and dim.min <= i <= dim.max:
                return dim.width
        return None
    for col, dim in a.column_dimensions.items():
        if dim.width:
            d(f"[{nome}] largura {col}", largura(a, col), largura(b, col))
    for row in range(1, max(a.max_row, b.max_row) + 1):
        for c in range(1, max(a.max_column, b.max_column) + 1):
            x, y = a.cell(row, c), b.cell(row, c)
            onde = f"[{nome}] {x.coordinate}"
            vx, vy = norm(x.value), norm(y.value)
            if vx in ("", None) and vy in ("", None):
                vx = vy = None
            d(onde + " valor", vx, vy)
            if vx is None and not x.has_style:
                continue
            n_cel += 1
            d(onde + " formato", x.number_format, y.number_format)
            d(onde + " fonte", (bool(x.font.b), bool(x.font.i), x.font.sz if x.font.sz != 11 else None,
                                (x.font.color.rgb if x.font.color is not None and x.font.color.type == "rgb" else None)),
              (bool(y.font.b), bool(y.font.i), y.font.sz if y.font.sz != 11 else None,
               (y.font.color.rgb if y.font.color is not None and y.font.color.type == "rgb" else None)))
            fx = x.fill.fgColor.rgb if x.fill and x.fill.fill_type == "solid" else None
            fy = y.fill.fgColor.rgb if y.fill and y.fill.fill_type == "solid" else None
            d(onde + " preenchimento", fx, fy)
            d(onde + " alinhamento", (x.alignment.horizontal, x.alignment.vertical, bool(x.alignment.wrap_text)),
              (y.alignment.horizontal, y.alignment.vertical, bool(y.alignment.wrap_text)))
            d(onde + " borda", tuple(getattr(x.border, s).style for s in ("left", "right", "top", "bottom")),
              tuple(getattr(y.border, s).style for s in ("left", "right", "top", "bottom")))

    def cfs(ws):
        out = []
        for rng, regras in ws.conditional_formatting._cf_rules.items():
            for r in regras:
                out.append((str(rng.sqref), r.type, tuple(r.formula or ()),
                            r.dxf.fill.fgColor.rgb if r.dxf and r.dxf.fill else None))
        return sorted(out)
    d(f"[{nome}] formatação condicional", cfs(a), cfs(b))


def graficos(arq):
    with zipfile.ZipFile(arq) as z:
        return [z.read(n).decode("utf-8") for n in sorted(z.namelist()) if n.startswith("xl/charts/chart")]


ga, gb = graficos(arq_py), graficos(arq_ts)
d("nº de gráficos", len(ga), len(gb))
for i, (x, y) in enumerate(zip(ga, gb), 1):
    d(f"gráfico {i} (XML)", x, y)


def ancoras(arq):
    import re
    with zipfile.ZipFile(arq) as z:
        return sorted(re.findall(r"<from>.*?</from><ext[^>]*>", z.read(n).decode("utf-8"))
                      for n in z.namelist() if n.startswith("xl/drawings/drawing"))


d("posição dos gráficos", ancoras(arq_py), ancoras(arq_ts))
for x in dif[:40]:
    print("DIFERENTE", x)
print(f"{len(dif)} diferença(s)" if dif else
      f"TUDO IGUAL — {len(wa.sheetnames)} abas, {n_cel} células com valor/estilo, {len(ga)} gráfico(s)")
